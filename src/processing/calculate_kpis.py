import json
import os
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from kpi_model import load_weights
from normalize_metrics import normalize_metrics

weights = load_weights()

OUTPUT_DIR = "data/processed"


def calculate_kpi(normalized: dict, weights: dict) -> float:
    total_weight = sum(weights.values())
    if not (0.99 <= total_weight <= 1.01):
        print(f"WARNING: Weights sum to {total_weight:.4f}, expected 1.0")

    weighted_sum = sum(
        normalized[metric] * weight
        for metric, weight in weights.items()
        if metric in normalized
    )
    return round(max(0.0, min(10.0, weighted_sum)), 4)


def build_dataset(data: dict, normalized: dict, kpi_score: float, weights: dict) -> dict:
    record = {
        "project":              data.get("project", "unknown"),
        "timestamp":            data.get("timestamp", ""),
        # raw metrics
        "raw_bugs":             data["metrics"]["bugs"],
        "raw_vulnerabilities":  data["metrics"]["vulnerabilities"],
        "raw_code_smells":      data["metrics"]["code_smells"],
        "raw_coverage":         data["metrics"]["coverage"],
        # normalized metrics (0-10)
        "norm_bugs":            normalized["bugs"],
        "norm_vulnerabilities": normalized["vulnerabilities"],
        "norm_code_smells":     normalized["code_smells"],
        "norm_coverage":        normalized["coverage"],
        # weights applied
        "weight_bugs":          weights["bugs"],
        "weight_vulnerabilities": weights["vulnerabilities"],
        "weight_code_smells":   weights["code_smells"],
        "weight_coverage":      weights["coverage"],
        # final KPI
        "kpi_score":            kpi_score,
    }
    return record


def save_csv(record: dict, path: str) -> None:
    df = pd.DataFrame([record])
    df.to_csv(path, index=False)
    print(f"CSV saved → {path}")


def save_xlsx(record: dict, path: str) -> None:
    df = pd.DataFrame([record])
    df.to_excel(path, index=False, sheet_name="KPI Report")

    wb = load_workbook(path)
    ws = wb["KPI Report"]

    # --- styles ---
    header_font    = Font(name="Arial", bold=True, color="FFFFFF", size=10)
    header_fill    = PatternFill("solid", start_color="1F4E79")   # dark blue
    section_fills  = {
        "raw":    PatternFill("solid", start_color="D9E1F2"),     # light blue
        "norm":   PatternFill("solid", start_color="E2EFDA"),     # light green
        "weight": PatternFill("solid", start_color="FFF2CC"),     # light yellow
        "kpi":    PatternFill("solid", start_color="FCE4D6"),     # light orange
    }
    kpi_font       = Font(name="Arial", bold=True, size=11, color="C00000")
    center         = Alignment(horizontal="center", vertical="center")
    thin           = Side(style="thin", color="BFBFBF")
    border         = Border(left=thin, right=thin, top=thin, bottom=thin)

    # column → section mapping
    col_section = {}
    for i, col in enumerate(df.columns, start=1):
        if col.startswith("raw_"):
            col_section[i] = "raw"
        elif col.startswith("norm_"):
            col_section[i] = "norm"
        elif col.startswith("weight_"):
            col_section[i] = "weight"
        elif col == "kpi_score":
            col_section[i] = "kpi"

    # header row styling
    for cell in ws[1]:
        cell.font      = header_font
        cell.fill      = header_fill
        cell.alignment = center
        cell.border    = border

    # data row styling
    for cell in ws[2]:
        col_idx = cell.column
        section = col_section.get(col_idx)
        if section:
            cell.fill = section_fills[section]
        if section == "kpi":
            cell.font = kpi_font
        cell.alignment = center
        cell.border    = border

    # auto column width
    for col_idx, col_cells in enumerate(ws.iter_cols(min_row=1, max_row=2), start=1):
        max_len = max(len(str(c.value or "")) for c in col_cells)
        ws.column_dimensions[get_column_letter(col_idx)].width = max_len + 4

    # freeze header
    ws.freeze_panes = "A2"

    wb.save(path)
    print(f"XLSX saved → {path}")


def formula_description():
    print("Loaded KPI weights:", weights)

    with open("data/raw/sonar_metrics.json", "r") as f:
        data = json.load(f)

    raw_metrics = data["metrics"]

    normalized = normalize_metrics(raw_metrics)
    print("Normalized metrics:", normalized)

    kpi_score = calculate_kpi(normalized, weights)
    print(f"KPI Score: {kpi_score} / 10")

    record = build_dataset(data, normalized, kpi_score, weights)

    os.makedirs(OUTPUT_DIR, exist_ok=True)
    save_csv(record,  os.path.join(OUTPUT_DIR, "kpi_results.csv"))
    save_xlsx(record, os.path.join(OUTPUT_DIR, "kpi_results.xlsx"))


if __name__ == "__main__":
    formula_description()